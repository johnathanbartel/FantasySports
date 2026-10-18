"""
Offline Draft Mode phase 4 through the real HTTP API on a throwaway SQLite file (offline_excel_smoke_test.db): the Excel
template (sheets and visibility, dropdown lists, protection, prefilled draft order, unique display names, auction budget
block), and the upload (session / version / format checks, exact-name player lookup, PlayerID path, every row-level error,
formulas never evaluated, preview diff, Replace all vs Fill empty cells only, entry_seq in row order, lands in review).

    .venv-api/Scripts/python.exe scripts/seed_offline_excel_smoke_test.py
"""

import io
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

db_path = Path(os.environ.get("OFFLINE_EXCEL_SMOKE_DB", "offline_excel_smoke_test.db"))
os.environ["DATABASE_URL"] = f"sqlite:///./{db_path.as_posix()}"
os.environ.pop("FSC_GRANT_ALL_PREMIUM", None)
if db_path.exists():
    db_path.unlink()

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from fantasy_sports.api.main import app
from fantasy_sports.platform import offline_excel as xl
from fantasy_sports.platform.db.base import Base
from fantasy_sports.platform.db.models.core import (
    DraftableAsset, League, LeagueUniverse, Player, RealTeam, Team, Universe, User,
)
from fantasy_sports.platform.db.models.offline_draft import OfflineDraftPick, OfflineDraftSession
from fantasy_sports.platform.db.session import SessionLocal, engine
from fantasy_sports.sports.hockey.models import Shot, ShotCredit  # noqa: F401

Base.metadata.create_all(bind=engine)
client = TestClient(app, follow_redirects=False)
db = SessionLocal()
checks = 0


def check(cond, msg):
    global checks
    checks += 1
    print(("ok    " if cond else "FAIL  ") + msg)
    if not cond:
        raise SystemExit(1)


# ---- seed ---------------------------------------------------------------------------------------------------------------------
users = {"commish": User(email="commish@example.com", display_name="commish"),
         "manager": User(email="manager@example.com", display_name="manager")}
db.add_all(users.values())
nhl = Universe(sport="hockey", level="pro", code="nhl", display_name="NHL")
db.add(nhl)
db.flush()
league = League(name="Excel League", commissioner_user_id=users["commish"].id, size=4, draft_rounds=3, draft_type="snake",
                series_length_days=7)
db.add(league)
db.flush()
db.add(LeagueUniverse(league_id=league.id, universe_id=nhl.id, defense_mode="group"))
owners = [users["commish"].id, users["manager"].id, None, None]
teams = [Team(league_id=league.id, name=n, owner_user_id=o, is_ai_owned=o is None, program_size="small")
         for n, o in zip(["Boston Bruins", "Team Two", "Team Three", "Team Four"], owners)]
bos = RealTeam(universe_id=nhl.id, name="Boston", external_ref="BOS")
tor = RealTeam(universe_id=nhl.id, name="Toronto", external_ref="TOR")
db.add_all(teams + [bos, tor])
db.flush()
NAMES = ["Connor McDavid", "Auston Matthews", "Ken Walker III", "Sam Reinhart", "Sam Reinhart", "Cale Makar", "Adam Fox",
         "Igor Shesterkin", "Nathan MacKinnon", "Leon Draisaitl", "Quinn Hughes", "Jack Eichel", "Mitch Marner",
         "Elias Pettersson", "Sidney Crosby", "Cher"]
POS = ["C", "C", "W", "W", "W", "D", "D", "G", "C", "C", "D", "C", "W", "C", "C", "W"]
assets = []
for i, (nm, pos) in enumerate(zip(NAMES, POS)):
    p = Player(universe_id=nhl.id, real_team_id=bos.id if i % 2 == 0 else tor.id, full_name=nm, position=pos)
    db.add(p)
    db.flush()
    a = DraftableAsset(universe_id=nhl.id, asset_type="player", player_id=p.id)
    db.add(a)
    assets.append(a)
db.commit()
LID = league.id
T = [t.id for t in teams]
A = [a.id for a in assets]


def login(user):
    resp = client.post("/sessions", json={"user_id": user.id})
    return {"Authorization": f"Bearer {resp.json()['token']}"}


H = {n: login(u) for n, u in users.items()}


def post(path, body=None, who="commish"):
    return client.post(path, json=body or {}, headers=H[who])


def upload(path, data, who="commish"):
    return client.post(path, content=data, headers={**H[who], "Content-Type": "application/octet-stream"})


def fill(data, edits, meta=None):
    """edits: {row: {header: value}} applied to the Picks sheet; meta: {key: value} for _Meta. Returns new xlsx bytes."""
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Picks"]
    head = {c.value: c.column for c in ws[1] if c.value}
    for row, vals in edits.items():
        for h, v in vals.items():
            ws.cell(row=row, column=head[h]).value = v
    if meta:
        for r in wb["_Meta"].iter_rows(min_row=1, max_col=2):
            if r[0].value in meta:
                r[1].value = meta[r[0].value]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ==== SNAKE TEMPLATE ============================================================================================================
S = post(f"/leagues/{LID}/offline-draft", {"format": "snake", "rounds": 3}).json()["id"]
base = f"/offline-drafts/{S}"
check(client.get(f"{base}/template", headers=H["manager"]).status_code == 403, "a manager can't download the template")
r = client.get(f"{base}/template", headers=H["commish"])
check(r.status_code == 200 and "spreadsheetml" in r.headers["content-type"], "the commissioner downloads an .xlsx")
tpl = r.content
wb = load_workbook(io.BytesIO(tpl))
check(wb.sheetnames == ["Instructions", "Picks", "Players", "Teams", "_Meta"], f"sheets: {wb.sheetnames}")
check(wb["Players"].sheet_state == "hidden" and wb["Teams"].sheet_state == "hidden" and wb["_Meta"].sheet_state == "veryHidden", "Players and Teams are hidden, _Meta is very hidden")
check(all(wb[n].protection.sheet for n in ("Picks", "Players", "Teams", "_Meta")), "Picks and the lookup sheets are protected")
ws = wb["Picks"]
check([c.value for c in ws[1]] == ["Round", "Pick", "Overall", "Team", "Player", "PlayerID", "Keeper"], "snake columns")
check(ws.max_row == 13, "exactly rounds x teams rows (12) plus the header")
check(ws["D2"].value == "Boston Bruins" and ws["D6"].value == "Team Four" and ws["C6"].value == 5, "Team is prefilled from the snake order (round 2 reverses)")
check(ws.column_dimensions["F"].hidden and ws["F2"].value.startswith("=IFERROR(INDEX(Players!$A:$A,MATCH(E2,Players!$B:$B,0))"), "PlayerID is a hidden INDEX/MATCH formula")
check(not ws["E2"].protection.locked and not ws["D2"].protection.locked and ws["A2"].protection.locked and ws["F2"].protection.locked, "only Team, Player and Keeper are editable")
dvs = {str(dv.sqref): dv.formula1 for dv in ws.data_validations.dataValidation}
check(dvs.get("D2:D13") == "=TeamNames" and dvs.get("E2:E13") == "=PlayerNames", "Team and Player have dropdown lists")
check(set(wb.defined_names.keys()) == {"PlayerNames", "TeamNames"}, "the lists are named ranges")
names = [r[1].value for r in wb["Players"].iter_rows(min_row=2)]
check(names == sorted(names) and len(names) == 16 and len(set(names)) == 16, "player names are unique and sorted")
check("Reinhart, Sam – BOS W" in names and "Reinhart, Sam – TOR W" in names, "format is Last, First – TEAM POS")
check("Walker III, Ken – BOS W" in names, "a suffix stays with the last name")
dup = [n for n in names if n.startswith("Reinhart")]
check(len(dup) == 2, "same-name players on different teams are distinguished")
meta = {r[0].value: r[1].value for r in wb["_Meta"].iter_rows(min_row=1, max_col=2)}
check(meta["session_id"] == S and meta["format"] == "snake" and meta["template_version"] == 1 and meta["settings_hash"], "_Meta carries the session, format, version and settings hash")
check(any("COUNTIF" in str(rule.formula) for rng in ws.conditional_formatting for rule in rng.rules), "duplicate players are highlighted")

# ==== SNAKE UPLOAD: errors ======================================================================================================
PN = {a.id: n for a, n in zip(assets, [None] * 16)}
db.expire_all()
name_of = xl.display_names(db, db.get(OfflineDraftSession, S))
nm = {aid: d["name"] for aid, d in name_of.items()}      # asset id -> dropdown text
good = {i + 2: {"Player": nm[A[i]]} for i in range(12)}   # row 2.. in board order
prev = lambda data, mode="replace_all": upload(f"{base}/upload/preview?mode={mode}", data)

check(prev(b"not a workbook").status_code == 422, "a non-xlsx upload is refused")
check(prev(b"PK" + b"\0" * (3 * 1024 * 1024)).json()["detail"]["message"].startswith("The file is too large"), "an over-limit file is refused")
check(upload(f"{base}/upload/preview", tpl, who="manager").status_code == 403, "a manager can't upload")
other = post(f"/leagues/{LID}/offline-draft", {"format": "snake"})   # refused: one open draft per league
bad = fill(tpl, good, meta={"session_id": 999})
check("different draft" in prev(bad).json()["detail"]["message"], "a template from another session is refused")
check("version" in prev(fill(tpl, good, meta={"template_version": 99})).json()["detail"]["message"], "an unsupported template version is refused")
check("auction" in prev(fill(tpl, good, meta={"format": "auction"})).json()["detail"]["message"], "a wrong format is refused")

edits = dict(good)
edits[4] = {"Player": "K Walker"}                                # not in the list
edits[6] = {"Player": nm[A[0]]}                                  # duplicate of row 2
edits[7] = {**good[7], "Team": "Nobody FC"}
edits[8] = {**good[8], "Round": 1, "Pick": 1}                               # duplicate cell with row 2
edits[9] = {**good[9], "Keeper": "X"}
edits[10] = {**good[10], "Overall": 99}
r = prev(fill(tpl, edits))
errs = r.json()["detail"]["errors"]
check(r.status_code == 422, "a workbook with bad rows is rejected as a whole")
check('Row 4: "K Walker" is not a valid player. Use the dropdown.' in errs, 'Row 4: "K Walker" is not a valid player. Use the dropdown.')
check(any(e.startswith("Row 6:") and "also on row 2" in e for e in errs), "a duplicate player names both rows")
check('Row 7: "Nobody FC" is not a valid team. Use the dropdown.' in errs, "an unknown team is named")
check(any(e.startswith("Row 8:") and "Round 1, Pick 1" in e for e in errs), "a duplicate board slot is caught")
check(any(e.startswith("Row 9:") and "Keeper" in e for e in errs), "a bad Keeper value is caught")
check(any(e.startswith("Row 10:") and "Overall" in e for e in errs), "an inconsistent Overall is caught")
check([int(e.split()[1].rstrip(":")) for e in errs if e.startswith("Row ")] == sorted(int(e.split()[1].rstrip(":")) for e in errs if e.startswith("Row ")), "errors are listed in row order")
# roster limit: four players on one team
edits = {i + 2: {"Player": nm[A[i]], "Team": "Team Two"} for i in range(4)}
check(any("Team Two has 4 players; the limit is 3" in e for e in prev(fill(tpl, edits)).json()["detail"]["errors"]), "a roster over the limit is refused")
# formulas are never evaluated; PlayerID path; mismatch
r = prev(fill(tpl, {2: {"Player": "=1+1"}, 3: {"Player": "=HYPERLINK(\"http://x\")"}}))
check(r.status_code == 200 and r.json()["rows"] == 0, "a formula typed into a cell is not evaluated or imported")
r = prev(fill(tpl, {2: {"PlayerID": A[3]}}))
check(r.status_code == 200 and r.json()["add"][0]["player"] == nm[A[3]], "a PlayerID alone resolves the player")
r = prev(fill(tpl, {2: {"PlayerID": A[3], "Player": nm[A[4]]}}))
check(r.status_code == 422 and "does not match its PlayerID" in r.json()["detail"]["errors"][0], "a Player that disagrees with its PlayerID is refused")
r = prev(fill(tpl, {2: {"PlayerID": 987654}}))
check(r.status_code == 422, "an unknown PlayerID with no name is refused")
check(db.query(OfflineDraftPick).count() == 0, "preview and rejected uploads wrote nothing")

# ==== SNAKE UPLOAD: preview + fill empty + replace all ==========================================================================
# two manual picks first: one the upload agrees with (cell 1.1), one it contradicts (cell 1.2)
check(post(f"{base}/picks", {"asset_id": A[0], "round": 1, "slot": 1}).status_code == 200, "manual pick 1.1 (agrees with the upload)")
check(post(f"{base}/picks", {"asset_id": A[15], "round": 1, "slot": 2}).status_code == 200, "manual pick 1.2 (contradicts the upload)")
full = fill(tpl, good)
p = prev(full, "fill_empty").json()
check(p["unchanged"] == 1 and len(p["add"]) == 10 and len(p["skipped"]) == 1 and p["skipped"][0]["why"] == "cell already filled", f"fill-empty preview: 1 unchanged, 10 to add, 1 skipped ({p['unchanged']}/{len(p['add'])}/{len(p['skipped'])})")
p = prev(full, "replace_all").json()
check(p["unchanged"] == 1 and len(p["add"]) == 11 and len(p["replace"]) == 1, "replace-all preview: the contradicting pick would be replaced")
check(upload(f"{base}/upload/commit?mode=nonsense", full).status_code == 400, "an unknown mode is refused")
r = upload(f"{base}/upload/commit?mode=fill_empty", full)
check(r.status_code == 200 and r.json()["result"]["added"] == 10, "fill-empty commits the empty cells only")
snap = r.json()
live = [x for x in snap["picks"] if not x["voided"]]
check(snap["status"] == "review", "an uploaded draft lands in review")
check(any(x["asset_id"] == A[15] and x["round"] == 1 and x["slot"] == 2 for x in live), "...and the manual pick the upload contradicted was kept")
up = sorted((x for x in live if x["source"] == "upload"), key=lambda x: x["entry_seq"])
check(len(up) == 10 and [x["overall"] for x in up] == sorted(x["overall"] for x in up), "uploaded picks take entry_seq in row order")
# replace all with a reshuffle
swap = {2: {"Player": nm[A[1]]}, 3: {"Player": nm[A[0]]}}
r = upload(f"{base}/upload/commit?mode=replace_all", fill(tpl, {**good, **swap}))
check(r.status_code == 200 and r.json()["result"]["voided"] >= 1, "replace-all voids what it replaces")
live = [x for x in r.json()["picks"] if not x["voided"]]
cell = {(x["round"], x["slot"]): x for x in live}
check(len(live) == 12 and cell[(1, 1)]["asset_id"] == A[1] and cell[(1, 2)]["asset_id"] == A[0], "the board now matches the upload exactly")
check(len({x["asset_id"] for x in live}) == 12, "no player is on the board twice")
check(sum(1 for x in r.json()["picks"] if x["voided"]) >= 2, "history keeps the replaced picks, voided")
check(post(f"{base}/finalize", {}).status_code == 200, "an uploaded draft finalizes as usual")

# template round-trip: a template with the current board prefilled
post(f"{base}/reopen")
t2 = load_workbook(io.BytesIO(client.get(f"{base}/template", headers=H["commish"]).content))["Picks"]
check(t2["E2"].value == nm[A[1]] and t2["E3"].value == nm[A[0]], "a later template is prefilled with what's already on the board")

# ==== AUCTION ===================================================================================================================
post(f"{base}/finalize", {})                         # still in review after reopen
from fantasy_sports.platform.db.models.core import Draft, DraftPick, RosterEntry
db.query(RosterEntry).delete()
db.query(DraftPick).delete()
db.query(Draft).delete()
db.commit()
S2 = post(f"/leagues/{LID}/offline-draft", {"format": "auction", "budget": 100, "min_bid": 2, "rounds": 3}).json()["id"]
ab = f"/offline-drafts/{S2}"
atpl = client.get(f"{ab}/template", headers=H["commish"]).content
wb = load_workbook(io.BytesIO(atpl))
ws = wb["Picks"]
check([c.value for c in ws[1]][:5] == ["Seq", "Team", "Player", "Price", "PlayerID"] and ws.max_row == 13, "auction columns, teams x roster rows")
check(ws["G1"].value == "Team" and ws["H2"].value == 100 and ws["I2"].value.startswith("=SUMIF") and ws["J2"].value == "=H2-I2", "the budget block has live formulas")
check(any(str(r.sqref).startswith("J2") for r in ws.conditional_formatting), "remaining budget turns red when negative")
db.expire_all()
nm = {aid: d["name"] for aid, d in xl.display_names(db, db.get(OfflineDraftSession, S2)).items()}
aprev = lambda d, mode="replace_all": upload(f"{ab}/upload/preview?mode={mode}", d)
r = aprev(fill(atpl, {2: {"Team": "Team Two", "Player": nm[A[0]], "Price": 1}}))
check(r.status_code == 422 and "Price must be a whole number of at least 2" in r.json()["detail"]["errors"][0], "a price under the minimum is refused")
r = aprev(fill(atpl, {2: {"Team": "Team Two", "Player": nm[A[0]], "Price": 2.5}}))
check(r.status_code == 422, "a fractional price is refused")
r = aprev(fill(atpl, {2: {"Team": "Team Two", "Player": nm[A[0]], "Price": 70}, 3: {"Team": "Team Two", "Player": nm[A[1]], "Price": 50}}))
check(r.status_code == 422 and any("Team Two spends 120 but its budget is 100" in e for e in r.json()["detail"]["errors"]), "a team over budget is refused")
r = aprev(fill(atpl, {2: {"Team": "Team Two", "Player": nm[A[0]]}}))
check(r.status_code == 422, "a sale with no price is refused")
good_a = {2: {"Team": "Team Two", "Player": nm[A[0]], "Price": 40}, 3: {"Team": "Team Three", "Player": nm[A[1]], "Price": 25},
          4: {"Team": "Team Two", "Player": nm[A[2]], "Price": 30}}
p = aprev(fill(atpl, good_a)).json()
check(len(p["add"]) == 3 and p["add"][0]["where"] == "$40", "auction preview lists the sales")
check(post(f"{ab}/sales", {"asset_id": A[1], "team_id": T[2], "price": 25}).status_code == 200, "a sale already logged by hand")
p = aprev(fill(atpl, good_a)).json()
check(p["unchanged"] == 1 and len(p["add"]) == 2, "an identical sale is 'unchanged'")
r = upload(f"{ab}/upload/commit?mode=replace_all", fill(atpl, good_a))
snap = r.json()
live = [x for x in snap["picks"] if not x["voided"]]
t2s = [t for t in snap["teams"] if t["team_id"] == T[1]][0]
check(r.status_code == 200 and snap["status"] == "review" and len(live) == 3 and t2s["spent"] == 70 and t2s["remaining"] == 30, "auction upload commits, budgets follow")
check([x["price"] for x in sorted(live, key=lambda x: x["entry_seq"]) if x["source"] == "upload"] == [40, 30], "new sales take entry_seq in row order")
# fill empty must respect budgets already spent
over = {2: {"Team": "Team Two", "Player": nm[A[5]], "Price": 50}}
r = aprev(fill(atpl, over), "fill_empty")
check(r.status_code == 422 and any("Team Two would spend 120" in e for e in r.json()["detail"]["errors"]), "fill-empty checks the combined budget with what is already sold")

print(f"\nAll {checks} offline-excel checks passed.")
db.close()
engine.dispose()
if db_path.exists():
    db_path.unlink()
